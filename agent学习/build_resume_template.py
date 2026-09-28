from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\13617\Desktop\agent学习\output\docx\南科大_AI优化科研实习简历模板.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

FONT_CN = "Microsoft YaHei"
FONT_EN = "Aptos"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(82, 88, 92)
LIGHT_GRAY = RGBColor(120, 125, 128)


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT_EN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), FONT_EN)
    rfonts.set(qn("w:hAnsi"), FONT_EN)
    rfonts.set(qn("w:eastAsia"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.52)
section.bottom_margin = Inches(0.50)
section.left_margin = Inches(0.70)
section.right_margin = Inches(0.70)
section.header_distance = Inches(0.2)
section.footer_distance = Inches(0.2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT_EN
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
normal.font.size = Pt(10.2)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(1.5)
normal.paragraph_format.line_spacing = 1.06

title_style = styles["Title"]
title_style.font.name = FONT_EN
title_style._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(1)
title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

heading = styles["Heading 1"]
heading.font.name = FONT_EN
heading._element.rPr.rFonts.set(qn("w:ascii"), FONT_EN)
heading._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_EN)
heading._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
heading.font.size = Pt(11.5)
heading.font.bold = True
heading.font.color.rgb = BLACK
heading.paragraph_format.space_before = Pt(5.5)
heading.paragraph_format.space_after = Pt(2)
heading.paragraph_format.keep_with_next = True

if "Resume Entry" not in styles:
    entry_style = styles.add_style("Resume Entry", WD_STYLE_TYPE.PARAGRAPH)
else:
    entry_style = styles["Resume Entry"]
entry_style.base_style = normal
entry_style.paragraph_format.space_before = Pt(1.5)
entry_style.paragraph_format.space_after = Pt(1)
entry_style.paragraph_format.keep_with_next = True

if "Resume Bullet" not in styles:
    bullet_style = styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
else:
    bullet_style = styles["Resume Bullet"]
bullet_style.base_style = normal
bullet_style.paragraph_format.left_indent = Inches(0.18)
bullet_style.paragraph_format.first_line_indent = Inches(-0.14)
bullet_style.paragraph_format.space_after = Pt(0.5)
bullet_style.paragraph_format.line_spacing = 1.02


def add_placeholder(paragraph, text):
    run = paragraph.add_run(text)
    set_run_font(run, color=LIGHT_GRAY, italic=True)
    return run


def add_heading(text):
    p = doc.add_paragraph(text, style="Heading 1")
    return p


def add_entry(left, right=None, secondary=None):
    p = doc.add_paragraph(style="Resume Entry")
    if right:
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.02), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(left)
    set_run_font(run, bold=True)
    if secondary:
        run2 = p.add_run(f"  {secondary}")
        set_run_font(run2, color=GRAY)
    if right:
        p.add_run("\t")
        rr = p.add_run(right)
        set_run_font(rr, color=GRAY)
    return p


def add_bullet(text=None, placeholder=None):
    p = doc.add_paragraph(style="Resume Bullet")
    r = p.add_run("• ")
    set_run_font(r, bold=True)
    if text:
        r2 = p.add_run(text)
        set_run_font(r2)
    if placeholder:
        add_placeholder(p, placeholder)
    return p


# Header
p = doc.add_paragraph(style="Title")
add_placeholder(p, "[姓名]")
direct_ppr = p._p.get_or_add_pPr()
direct_border = direct_ppr.find(qn("w:pBdr"))
if direct_border is not None:
    direct_ppr.remove(direct_border)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("南方科技大学本科生  |  AI＋优化科研实习申请")
set_run_font(r, size=10.8, bold=True, color=GRAY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(3)
add_placeholder(p, "[手机]  |  [邮箱]  |  [微信]  |  [GitHub链接 可选]")

# Education
add_heading("教育背景")
add_entry("南方科技大学", "[入学年月] 至今", "[当前院系]  [年级]")
add_bullet(placeholder="拟选择专业：[数据科学与大数据技术或实际专业]")
add_bullet(placeholder="GPA：[X.XX/4.00]；排名：[X/X]（无优势可删除）")
add_bullet(placeholder="相关课程：[程序设计基础]、[线性代数]、[数学分析]、[概率论]（可附高分）")

# Awards
add_heading("竞赛与荣誉")
add_entry("全国青少年信息学奥林匹克竞赛 NOI 铜牌", "[年份]")
add_bullet(text="长期使用C++进行算法训练，较擅长数据结构与构造类问题。")
add_bullet(placeholder="[其他含金量较高的数学 信息学或英语奖项 没有则删除]")

# Projects
add_heading("项目与实践经历")
add_entry("信息学竞赛算法训练", "[起止年份]")
add_bullet(text="独立完成算法设计、复杂度分析、边界构造与代码调试。")
add_bullet(text="系统训练数据结构、图论、动态规划和构造算法，能够针对错误结果设计测试数据。")

add_entry("[机器学习 大模型或个人项目名称]", "[时间]")
add_bullet(placeholder="问题与方法：[解决什么问题 使用什么方法 你具体负责什么]")
add_bullet(placeholder="实验与结果：[数据集 基线 指标 运行效率 收敛结果或代码链接]")

# Skills
add_heading("技术能力")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0.5)
for label, value in [
    ("编程", "C++熟练；Python[基础或熟悉，按实际填写]"),
    ("算法", "数据结构、图论、动态规划、构造算法、复杂度分析"),
    ("机器学习", "[PyTorch 深度学习 强化学习等；尚在学习就明确写正在学习]"),
    ("工具", "[Git Linux LaTeX；只保留真实使用过的]"),
]:
    run = p.add_run(f"{label}：")
    set_run_font(run, bold=True)
    run2 = p.add_run(value + "    ")
    set_run_font(run2, color=GRAY if "[" in value else BLACK)

# Research interest
add_heading("研究兴趣与申请动机")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
text = (
    "对AI＋优化、强化学习、多智能体协作与LLM决策感兴趣，尤其希望研究agent面对未知交互对象时，"
    "如何根据行为反馈更新判断并优化后续策略。我偏好先提出可检验的假设，再通过代码与数值实验验证，"
    "希望接受从选题、方法实现到实验分析和论文写作的完整科研训练。"
)
r = p.add_run(text)
set_run_font(r)

# Availability
add_heading("时间安排与其他兴趣")
add_bullet(placeholder="学期中每周可投入：[X]小时；寒暑假：[可以或无法]参加宁波线下实习。")
add_bullet(text="其他学术兴趣：心理学、认知科学与哲学；", placeholder="个人爱好：[请填写]。")

# Small editing note
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(3)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("填写提示：删除所有方括号提示和不适用条目，最终控制在一页。")
set_run_font(r, size=8.5, color=LIGHT_GRAY, italic=True)

# Prevent accidental metadata leakage
doc.core_properties.title = "南方科技大学 AI优化科研实习简历模板"
doc.core_properties.subject = "本科科研实习申请简历"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""

doc.save(OUT)
print(OUT)
